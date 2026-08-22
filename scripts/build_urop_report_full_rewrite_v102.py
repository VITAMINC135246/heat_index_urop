from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "outputs" / "reports" / "full_rewrite_2026-08-05"
ASSET_DIR = OUT_DIR / "_build_assets"
OUT = OUT_DIR / "UROP_report_full_rewrite_v102_2026-08-05.docx"

NAVY = RGBColor(31, 77, 120)
BLUE = RGBColor(46, 116, 181)
GRAY = RGBColor(92, 99, 110)
LIGHT = "F4F6F9"
WHITE = RGBColor(255, 255, 255)


def font(run, size=11, bold=False, italic=False, color=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd")) or OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    if shd.getparent() is None:
        tc_pr.append(shd)


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}")) or OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")
        if node.getparent() is None:
            tc_mar.append(node)


def table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(total)); tbl_w.set(qn("w:type"), "dxa")
    ind = tbl_pr.find(qn("w:tblInd")) or OxmlElement("w:tblInd")
    ind.set(qn("w:w"), "120"); ind.set(qn("w:type"), "dxa")
    if ind.getparent() is None: tbl_pr.append(ind)
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol"); col.set(qn("w:w"), str(width)); grid.append(col)
    for row_i, row in enumerate(table.rows):
        if row_i == 0:
            tr_pr = row._tr.get_or_add_trPr(); rep = OxmlElement("w:tblHeader"); rep.set(qn("w:val"), "true"); tr_pr.append(rep)
        cant = OxmlElement("w:cantSplit"); row._tr.get_or_add_trPr().append(cant)
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(width)); tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(cell)


def add_table(doc, headers, rows, widths, caption):
    p = doc.add_paragraph(style="Caption")
    p.paragraph_format.keep_with_next = True
    font(p.add_run(caption), size=9.5, italic=True, color=GRAY)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        shade(table.rows[0].cells[i], LIGHT)
        para = table.rows[0].cells[i].paragraphs[0]; para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font(para.add_run(h), size=9, bold=True, color=NAVY)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            para = cells[i].paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT if len(str(value)) > 18 else WD_ALIGN_PARAGRAPH.CENTER
            font(para.add_run(str(value)), size=8.7)
    table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_captioned_figure(doc, path, caption, alt, width=6.35):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
    run = p.add_run(); run.add_picture(str(path), width=Inches(width))
    shape = doc.inline_shapes[-1]
    shape._inline.docPr.set("descr", alt)
    cp = doc.add_paragraph(style="Caption"); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.keep_together = True
    font(cp.add_run(caption), size=9.2, italic=True, color=GRAY)


def add_body(doc, text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    font(p.add_run(text))
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet"); font(p.add_run(item))


def add_heading(doc, text, level):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(paragraph.add_run("Page "), size=9, color=GRAY)
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def workflow_diagram(path):
    W, H = 2200, 940
    img = Image.new("RGB", (W, H), "white"); d = ImageDraw.Draw(img)
    regular = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 28)
    bold = ImageFont.truetype(r"C:\Windows\Fonts\arialbd.ttf", 34)
    small = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 23)
    boxes = [
        (60, 125, 360, 330, "User-selected\ninputs", "Visible, thermal,\nTAT3 or matrix"),
        (430, 125, 730, 330, "Part A", "Validation, identity,\ntime and dimensions"),
        (800, 125, 1100, 330, "Part B", "Correspondence evidence\nand explicit review"),
        (1170, 45, 1530, 265, "Normal route", "Accepted V/T ROI and\nreviewed physical cover"),
        (1170, 360, 1530, 580, "Polygon route", "Thermal target selected\nwhen V/T is unusable"),
        (1600, 125, 1900, 330, "Part D", "Temperature extraction\nand radiometric QA"),
        (1600, 505, 1900, 710, "Canonical store", "Versioned per-capture\nresult and cache"),
        (1960, 300, 2160, 530, "Part E", "Delta T, spectra,\nstatistics and maps"),
    ]
    for x1,y1,x2,y2,title,sub in boxes:
        d.rounded_rectangle((x1,y1,x2,y2), radius=24, fill="#F4F6F9", outline="#2E74B5", width=5)
        d.multiline_text(((x1+x2)//2, y1+22), title, font=bold, fill="#1F4D78", anchor="ma", align="center", spacing=3)
        d.multiline_text(((x1+x2)//2, y1+112), sub, font=small, fill="#40464F", anchor="ma", align="center", spacing=5)
    def arrow(a,b):
        d.line((a,b), fill="#2E74B5", width=7); x,y=b; d.polygon([(x,y),(x-22,y-13),(x-22,y+13)], fill="#2E74B5")
    arrow((360,228),(430,228)); arrow((730,228),(800,228)); arrow((1100,200),(1170,155)); arrow((1100,260),(1170,470))
    arrow((1530,155),(1600,210)); arrow((1530,470),(1600,270)); arrow((1750,330),(1750,505)); arrow((1900,605),(1960,440))
    d.rounded_rectangle((630, 770, 1880, 900), radius=24, fill="#EAF2F8", outline="#1F4D78", width=4)
    d.text((1255,800), "Optional temporal analysis: explicit grouping of compatible cached captures;\npixelwise change requires cross-capture registration", font=regular, fill="#1F4D78", anchor="ma", align="center")
    arrow((1750,710),(1750,770))
    img.save(path)


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True); ASSET_DIR.mkdir(parents=True, exist_ok=True)
    workflow = ASSET_DIR / "workflow_architecture_v102.png"; workflow_diagram(workflow)
    doc = Document(); sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11); sec.top_margin = sec.bottom_margin = Inches(1); sec.left_margin = sec.right_margin = Inches(1)
    sec.different_first_page_header_footer = True
    hp = sec.header.paragraphs[0]; font(hp.add_run("UROP Research Report | Full rewrite v102"), size=9, color=GRAY)
    page_field(sec.footer.paragraphs[0])
    styles = doc.styles
    normal = styles["Normal"]; normal.font.name = "Calibri"; normal.font.size = Pt(11); normal.paragraph_format.space_after = Pt(8); normal.paragraph_format.line_spacing = 1.333
    for name,size,color,before,after in (("Heading 1",16,BLUE,18,10),("Heading 2",13,BLUE,12,6),("Heading 3",12,NAVY,8,4)):
        st=styles[name]; st.font.name="Calibri"; st.font.size=Pt(size); st.font.color.rgb=color; st.font.bold=True; st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after)
    styles["Caption"].font.name="Calibri"; styles["Caption"].font.size=Pt(9.2)

    for _ in range(4): doc.add_paragraph()
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run("UNDERGRADUATE RESEARCH REPORT"),12,True,color=GRAY)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(18); font(p.add_run("Development and Validation of a User-Operated UAV Thermal Analysis Workflow"),27,True,color=NAVY)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run("Surface-temperature, delta-temperature, spatial and temporal assessment with HKUST validation cases"),14,italic=True,color=BLUE)
    doc.add_paragraph();
    for label,value in (("Student","[Student Name] ([Student ID])"),("Supervisor","[Supervisor Name]"),("Department","[Department]"),("Institution","Hong Kong University of Science and Technology"),("Draft","Full rewrite v102 - 5 August 2026")):
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run(f"{label}: "),10.5,True,color=NAVY); font(p.add_run(value),10.5)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(22); font(p.add_run("Independent complete rewrite; not the section-by-section v2 draft"),9.5,italic=True,color=GRAY)
    doc.add_page_break()

    add_heading(doc,"Abstract",1)
    add_body(doc,"This project developed and validated a user-operated workflow for converting paired unmanned-aerial-vehicle visible and radiometric thermal imagery into auditable surface-temperature, surface-to-ambient temperature-difference (Delta T), spatial and temporal results. The system was designed as a reusable analytical tool rather than a pipeline restricted to one fixed dataset. A user may select one image group or a discoverable dataset, review visible-thermal correspondence, continue through a normal visible-image surface-cover route or a thermal-polygon fallback, extract or ingest a quality-controlled native temperature grid, and generate versioned canonical results. Accepted per-capture results are retained for later reuse and may enter temporal analysis only after target identity, region-of-interest comparability, provenance and quality compatibility are explicitly confirmed.")
    add_body(doc,"Two validation cases demonstrate the implemented boundary. Five HKUST visible-thermal pairs validated the normal route and produced a 1,638,400-row full-pixel dataset. Deterministic spatial thinning supported distributional analysis without treating all neighbouring pixels as independent replicates. Full-population median Delta T values were 2.47 degC for tree vegetation, 5.71 degC for concrete pavement, 6.65 degC for roof, 12.36 degC for grass or low vegetation, and 14.15 degC for bare soil; these are descriptive winter-afternoon case results, not universal material rankings. A separate football-field case validated the polygon and temporal routes using three captures at 09:11, 14:08 and 17:04. All three were eligible for target-level comparison, while pixelwise change was correctly withheld because cross-capture registration had not been established.")
    add_body(doc,"The present system measures apparent surface temperature and Delta T relative to an image-specific ambient parameter. It does not calculate the Hong Kong Heat Index, estimate physiological heat strain or issue work-rest instructions. Its contribution is an extensible, provenance-preserving foundation for future heat-risk decision support after independent meteorological observations, radiometric calibration, expanded sampling and an accepted heat-stress formulation are incorporated.")
    add_body(doc,"Keywords: UAV thermal remote sensing; apparent surface temperature; delta temperature; reproducible workflow; spatial analysis; temporal analysis; occupational heat-risk decision support")

    add_heading(doc,"Contents",1)
    for item in ["1. Introduction","2. Literature Review","3. Methodology","4. Results","5. Discussion","6. Conclusion and Future Work","References","Appendices"]: add_body(doc,item)

    add_heading(doc,"1. Introduction",1)
    add_heading(doc,"1.1 Background",2)
    add_body(doc,"Urban surfaces differ in albedo, emissivity, moisture availability, roughness and heat-storage capacity, creating fine-scale thermal heterogeneity. Remotely sensed surface temperature is related to, but distinct from, near-surface air temperature and human heat exposure (Oke, 1982; Voogt & Oke, 2003). UAV systems can resolve individual roofs, pavements, vegetation and target areas that are obscured at coarser satellite resolution, but their apparent precision creates methodological risks when radiometry, cross-sensor alignment, label provenance and spatial dependence are not made explicit.")
    add_heading(doc,"1.2 Problem and research contribution",2)
    add_body(doc,"The project began under the broad theme of urban heat-index assessment. Implementation showed that the scientifically supportable current output is apparent surface temperature and surface-to-ambient Delta T, not a standard heat index. The central contribution is therefore a reusable analytical workflow that connects user-selected inputs, manual review, canonical storage, statistical analysis, mapping and guarded temporal comparison. The five-image pilot is a benchmark for the normal route; it is not the boundary of the system. A football-field case separately tests target selection and repeated-capture analysis.")
    add_heading(doc,"1.3 Long-term application pathway",2)
    add_body(doc,"A future application is site-specific heat-risk decision support for outdoor work, sport and campus operations. Surface maps could help identify locations and periods that merit closer monitoring, evaluate shade or surface interventions, and guide sensor placement. They cannot by themselves determine safe work-rest schedules. Hong Kong's Heat Stress at Work Warning is based on the Hong Kong Heat Index and considers meteorological heat stress at a population level; practical work arrangements additionally depend on workload and other risk factors (Hong Kong Labour Department, 2026). Any future advisory layer must therefore integrate validated air temperature, humidity, wind and radiation with an accepted occupational framework rather than applying an arbitrary Delta T threshold.")
    add_heading(doc,"1.4 Aim, objectives and research questions",2)
    add_body(doc,"The aim was to develop and validate a user-operated UAV thermal-analysis workflow that preserves scientific provenance from input selection to reusable spatial and temporal results.")
    add_bullets(doc,["RQ1. How can visible imagery, thermal imagery, LUHK context, reviewed physical cover, radiometric parameters and ambient metadata be integrated through auditable normal and fallback routes?","RQ2. Can accepted per-capture results be stored and reused while preserving target eligibility, measurement type, provenance, schema identity, cache validity and temporal compatibility?","RQ3. What functional and scientific evidence is produced by the five-image normal-route benchmark and the three-capture football-field polygon/temporal case, and what limitations constrain interpretation?"])

    add_heading(doc,"2. Literature Review",1)
    add_heading(doc,"2.1 Surface temperature, heat index and scale",2)
    add_body(doc,"Thermal remote sensing observes radiometric surface behaviour, whereas atmospheric heat exposure depends on air temperature, humidity, airflow and radiation. This distinction prevents a hot roof pixel from being interpreted directly as pedestrian heat stress. Land-surface temperature relationships with vegetation and imperviousness also vary with time, season, moisture and spatial scale (Weng, Lu, & Schubring, 2004). The report therefore treats heat-index assessment as a future integration target rather than a completed output.")
    add_heading(doc,"2.2 UAV radiometry and visible-thermal correspondence",2)
    add_body(doc,"UAV photogrammetry offers flexible, very-high-resolution observation, but geometric accuracy depends on sensor calibration, viewing geometry, registration and ground control (Colomina & Molina, 2014). Visible and thermal cameras differ in optics, resolution and field of view. Thermal recovery additionally depends on emissivity, reflected apparent temperature, atmospheric conditions and object distance. Automated correspondence scores and successful SDK execution are useful evidence, not independent proof of spatial or radiometric truth.")
    add_heading(doc,"2.3 Land use, physical cover and spatial dependence",2)
    add_body(doc,"Planning-scale land use and material-scale physical cover answer different questions. LUHK provides official broad context, while reviewed visible-image masks describe the physical surface associated with thermal measurements. Raster pixels are also spatially dependent; millions of pixels do not constitute millions of independent experiments. This motivates explicit provenance, spatial thinning for exploratory comparisons and greater emphasis on distributions, image coverage and effect sizes than on small p-values alone.")

    add_heading(doc,"3. Methodology",1)
    add_heading(doc,"3.1 Development and validation design",2)
    add_body(doc,"A computational-method development and validation design was used. Workflow version v0.3.2 (processing identity heat-index-urop-0.3.2; canonical schema 0.2.0) was evaluated through tracked regression tests, controlled local acceptance and two real-data validation cases. Methodological decisions were encoded as review gates and explicit unavailable states rather than inferred from filenames, timestamps or spatial proximity.")
    add_heading(doc,"3.2 Inputs and validation cases",2)
    add_body(doc,"The general workflow accepts user-selected visible and thermal inputs, TAT3 reports or audited temperature matrices, optional official LUHK context, and reviewed target or surface-cover information. Native thermal dimensions are derived from each accepted temperature source. The five-image benchmark used five nadir DJI M4T pairs acquired at HKUST on 7 January 2026 and 512 by 640 temperature matrices. The second case used three football-field thermal captures on 2 February 2026 with user-accepted target polygons and image-specific ambient parameters.")
    add_heading(doc,"3.3 User workflow and A-E architecture",2)
    add_captioned_figure(doc,workflow,"Figure 1. Current user-operated A-E architecture, persistent canonical storage and optional temporal branch.","Flow diagram showing user-selected inputs, Parts A to E, normal and polygon routes, canonical storage, and temporal analysis.")
    add_body(doc,"The ordinary-user launcher accepts paths and review decisions interactively. Part A validates inputs; Part B presents correspondence evidence and requires an explicit decision; Part C either reviews physical cover on an accepted visible ROI or records a target polygon on the thermal image; Part D produces or audits temperatures; and Part E creates eligible statistics, spectra, maps and a readable result summary. Cancelled or rejected states cannot silently become successful results.")
    add_heading(doc,"3.4 Part A: validation, identity and pairing",2)
    add_body(doc,"Visible and thermal files are paired using capture identity and metadata rather than directory order alone. The workflow records missing, duplicate and unmatched states, validates timestamps and spatial metadata, derives native dimensions and checks compatibility between temperature grids and any masks. These checks establish the measurement identity required by later cache and temporal decisions.")
    add_heading(doc,"3.5 Part B: correspondence and alignment",2)
    add_body(doc,"Metadata and camera field-of-view assumptions initialise the thermal ROI within the visible image. Phase correlation, enhanced correlation coefficient optimisation, feature matching and local grid-search refinement may provide candidate evidence. Automatic scores are not final acceptance. The user reviews structural overlays and either accepts full visible support, rejects the correspondence and continues through the thermal-polygon route, or cancels the group. Manual ground-control-point refinement remains available when stable features can be identified.")
    add_heading(doc,"3.6 Part C: physical cover, LUHK and target scope",2)
    add_body(doc,"In the normal route, candidate visible-image segments are manually reviewed and mapped to the native thermal grid. LUHK 2024 remains a read-only broad land-use layer and cannot substitute for physical-cover review. Shadow is a separate binary property. In the fallback route, the user draws and accepts a target polygon on the thermal image; exterior pixels remain auditable but are target-false, label-unknown and ineligible for target statistics. User-supplied LUHK for a polygon is target-scoped context and is not represented as an official raster lookup.")
    add_heading(doc,"3.7 Part D: temperature extraction and radiometric QA",2)
    add_body(doc,"The active DJI route invokes dji_irp.exe in measure/float32 mode using per-image TAT3-associated distance, emissivity, humidity, ambient and reflected-temperature fields. Missing ambient preserves temperature but disables Delta T. The five-image case used emissivity 0.95 and humidity 50%, with reflected apparent temperature equal to the image-specific ambient parameter. These are case settings, not universal defaults. QA checks dimensional compatibility, finite and non-constant data, suspicious preview-like values, sub-zero pixels and apparent-temperature extremes. Unusual values are retained unless a prespecified scientific exclusion rule is applied.")
    add_heading(doc,"3.8 Part E: measurement types, Delta T and eligibility",2)
    add_body(doc,"The canonical schema supports full thermal pixels, polygon-selected thermal pixels, TAT3 points and TAT3 regions. Measurement types and sources remain separate unless an explicitly labelled sensitivity analysis requests pooling. For image i and thermal pixel p, Delta T was defined as:")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run("Delta T(i,p) = Tsurface(i,p) - Tambient(i)"),12,True,color=NAVY)
    add_body(doc,"The image-level ambient value is not imputed when missing. Positive Delta T indicates an apparent surface warmer than the recorded ambient parameter; it does not control solar radiation, wind, humidity, emissivity, view angle, shadow or thermal lag.")
    add_heading(doc,"3.9 Sampling and statistical analysis",2)
    add_body(doc,"The five-image case preserved all 1,638,400 finite pixels for descriptive summaries. Formal spectra and exploratory tests used deterministic spatial thinning: each 8 by 8 pixel tile contributed at most one original pixel, with no tile averaging. The primary seed was 20260715 and 20 secondary seeds assessed stability. SciPy gaussian_kde with Scott's bandwidth described distributions. Kruskal-Wallis and Mann-Whitney U tests with multiple-testing adjustment provided exploratory support; rank-biserial correlation and median differences described effect magnitude. Full-population summaries and sampled comparisons were reported separately.")
    add_heading(doc,"3.10 Canonical storage, cache and run isolation",2)
    add_body(doc,"Each accepted capture produces a versioned canonical native-grid result containing target, label-known, eligibility, source, provenance, temperature, ambient and QA fields. Persistent per-image results may be reused, while each analysis run writes its own run-scoped tables, figures and USER_RESULTS summary. Cache identity includes source and artifact hashes, review decisions, LUHK dependencies, polygon state, temperature and ambient definitions, capture time, configuration and implementation signatures. A relevant change invalidates affected downstream reuse.")
    add_heading(doc,"3.11 Spatial and temporal analysis",2)
    add_body(doc,"Spatial outputs preserve native dimensions and explicitly display unavailable layers. Temporal analysis is opt-in. One accepted capture of one physical target is one temporal observation; pixels within a capture are spatial measurements. Users must confirm same location, target identity and comparable ROIs. Capture-level mean, median, quantiles and extrema are calculated before cross-time comparison. Pixelwise range maps require compatible grids plus a documented cross-capture registration method and stable registration identifier. Missing periods are not interpolated, and partial observation windows are not described as complete daily cycles.")
    add_heading(doc,"3.12 Quality assurance and software validation",2)
    add_body(doc,"QA spans pairing, correspondence review, mask acceptance, temperature-grid compatibility, provenance, eligibility and output completeness. The accepted Phase 1 scientific baseline is frozen in version control. On 4 August 2026, the Windows CPython 3.12.13 reference environment passed dependency checking and 111 non-local regression tests; six controlled-local tests were deselected because they require private, proprietary or manually reviewed assets. Remote GitHub Actions validation remained pending when this draft was prepared.")

    doc.add_page_break()
    add_heading(doc,"4. Results",1)
    add_heading(doc,"4.1 Functional validation of the reusable workflow",2)
    add_table(doc,["Validation case","Route","Outcome","Interpretive boundary"],[["Five HKUST pairs","Normal V/T route","Five accepted/cached captures completed canonical, Part E and spatial outputs","Benchmark, not the system boundary"],["Football field","Thermal-polygon route","Accepted target-only statistics and maps; exterior excluded","User-defined target context"],["Three football-field captures","Explicit temporal group","3/3 eligible for target-level series","No pixelwise comparison without registration"]],[2000,1700,3100,2560],"Table 1. Validation cases used to test distinct workflow routes.")
    add_body(doc,"The acceptance runs demonstrated that the workflow is not fixed to the historical five-image dataset. It can route one group or a discoverable dataset, retain accepted per-capture results, combine compatible results in a new run and withhold outputs when scientific gates are unmet.")
    add_heading(doc,"4.2 Five-image benchmark completion and radiometric QA",2)
    add_body(doc,"All five normal-route pairs produced compatible 512 by 640 finite temperature grids, yielding 1,638,400 full-population pixel records. Mean apparent temperature ranged from 14.28 to 16.31 degC. The observed minimum was -29.15 degC and the maximum was 48.89 degC. Extraction succeeded structurally, but all five matrices contained values requiring physical plausibility review; successful software execution was therefore not treated as complete radiometric validation.")
    add_heading(doc,"4.3 Surface-cover Delta T distributions",2)
    rows=[["Tree vegetation","696,329","5","2.82","2.47","2.24"],["Roof","569,601","5","7.00","6.65","10.33"],["Concrete pavement","322,856","5","7.10","5.71","7.77"],["Grass / low vegetation","39,392","5","11.21","12.36","4.62"],["Bare soil","10,222","2","12.46","14.15","5.04"]]
    add_table(doc,["Physical cover","Full pixels","Images","Mean Delta T","Median Delta T","IQR"],rows,[2300,1300,800,1700,1700,1560],"Table 2. Full-population Delta T summaries for the five-image benchmark (degC).")
    spectrum=ROOT/"outputs/part_e/figures/spectrum/fig02_pixel_delta_t_spectrum_by_surface_cover_facets.png"
    add_captioned_figure(doc,spectrum,"Figure 2. Spatially thinned pixel-level Delta T density spectra by reviewed physical cover.","Five density plots showing sampled Delta T distributions for bare soil, concrete pavement, grass or low vegetation, roof, and tree vegetation.")
    add_body(doc,"The spatially thinned surface-cover sample contained 27,074 pixels. The global Kruskal-Wallis statistic was 3643.39. Concrete pavement and roof were not clearly separated after adjustment (p = 0.185862), consistent with strongly overlapping spectra. Sample-based rank-biserial correlations were 0.46 for concrete pavement versus tree vegetation, 0.33 for roof versus tree vegetation and 0.79 for grass/low vegetation versus tree vegetation. These values describe the sampled comparison; the medians in Table 2 describe the full pixel population.")
    add_heading(doc,"4.4 Spatial evidence",2)
    spatial=ROOT/"outputs/part_e/figures/spatial_maps/DJI_20260107143344_0009_combined_qa_panel.png"
    add_captioned_figure(doc,spatial,"Figure 3. Full-population native-grid QA panel for benchmark image 0009.","Maps of temperature, Delta T, LUHK code, physical cover and shadow for one benchmark image.")
    add_body(doc,"Native-grid panels showed coherent within-image thermal variation and demonstrated why LUHK and physical cover must remain separate. The example LUHK context is comparatively coarse, whereas the reviewed mask distinguishes multiple materials. The benchmark contained no valid shadow-present pixels, so a cover-by-shadow contrast was correctly recorded as not estimable.")
    add_heading(doc,"4.5 Football-field polygon and temporal case",2)
    temporal_rows=[["09:11:28","34,724","10.8","21.77","10.97"],["14:08:15","27,330","11.8","38.79","26.99"],["17:04:53","20,746","15.8","18.43","2.63"]]
    add_table(doc,["Local time","Eligible target pixels","Ambient","ROI mean T","ROI mean Delta T"],temporal_rows,[1600,2200,1300,2100,2160],"Table 3. Three-capture football-field target summaries (degC).")
    tfig=ROOT/"outputs/runs/v0_3_user_acceptance/part_e/schema_0_2/run_20260722T053855Z/temporal/hkust-football-field-20260202/figures/target_delta_t_vs_local_time.png"
    add_captioned_figure(doc,tfig,"Figure 4. Football-field target Delta T across the sampled observation window.","Three time points showing football-field target Delta T summaries at morning, afternoon and late afternoon.")
    add_body(doc,"All three captures were eligible for an observed target-level series. The maximum ROI-mean apparent temperature occurred at 14:08 (38.79 degC) and the minimum at 17:04 (18.43 degC), an observed difference of 20.36 degC. ROI-mean Delta T ranged from 2.63 to 26.99 degC, an observed difference of 24.36 degC. These are sampled-window results, not a complete daily range. The user confirmed a comparable physical target, but polygon boundaries varied and cross-capture registration was not established; consequently, pixelwise temporal change was unavailable by design.")
    add_heading(doc,"4.6 Software assurance",2)
    add_body(doc,"The local blocking regression suite passed 111 tests with six controlled-local tests excluded from the clean CI boundary. Canonical schema, target eligibility, routing, cache behaviour, Part E statistics, spatial figures and temporal compatibility were covered by automated tests. Proprietary SDK execution, interactive review, private real inputs and final rendered Excel review remain controlled-local activities. Remote CI had not yet run because the Phase 2 workflow was still local.")

    add_heading(doc,"5. Discussion",1)
    add_heading(doc,"5.1 Methodological contribution",2)
    add_body(doc,"The strongest contribution is not a fixed material ranking but an analytical system that preserves meaning across stages. It distinguishes preview intensity from radiometric temperature, official land use from physical cover, target pixels from audit-only exterior pixels, spatial measurements from temporal observations, persistent accepted captures from run-scoped outputs, and unavailable evidence from inferred data. These distinctions make later extension safer than an ad hoc collection of scripts or spreadsheets.")
    add_heading(doc,"5.2 Interpretation of benchmark patterns",2)
    add_body(doc,"Tree vegetation had the lowest and narrowest full-population Delta T distribution in the benchmark, while grass/low vegetation and bare soil were warmer. The contrast is physically plausible under differences in shade, evapotranspiration, moisture and exposure, but cover is confounded with image identity, location, orientation and thermal history. Bare soil occurred in only two images, and concrete pavement and roof overlapped strongly. The results therefore motivate hypotheses and demonstrate the analysis rather than establish universal causal material effects.")
    add_heading(doc,"5.3 Scientific and practical value",2)
    add_body(doc,"Academically, the workflow offers a provenance-aware way to study fine-scale surface thermal heterogeneity while limiting common forms of pseudoreplication. Practically, it can support target screening, repeated site monitoring, comparison of shade or surface interventions, and identification of locations requiring denser meteorological observation. The persistent canonical store is important because a capture accepted today can be reused in a later explicitly defined temporal group without treating unrelated directory contents as a time series.")
    add_heading(doc,"5.4 Path toward outdoor-worker heat-risk support",2)
    add_body(doc,"The football-field case illustrates a potential operational scenario: a target may appear much warmer relative to its recorded ambient parameter at one sampled time than at another. This evidence could prompt further measurement or operational review, but it is not itself a work-rest trigger. A defensible advisory system would combine calibrated surface information with the Hong Kong Heat Index or another accepted occupational measure, local meteorological sensors, workload and clothing, exposure duration, acclimatisation and workplace controls. The current tool should be viewed as a spatial evidence layer within that future system.")
    add_heading(doc,"5.5 Limitations",2)
    add_bullets(doc,["The five-image benchmark is one short winter-afternoon transect and cannot represent HKUST or Hong Kong.","Football-field results contain only three captures, no night observation and varying accepted polygon boundaries.","Image-level ambient values are TAT3-associated parameters rather than independently validated meteorological air temperatures.","Fixed emissivity and other radiometric assumptions may bias material comparisons; sub-zero and extreme pixels remain under review.","Visible-thermal alignment and approximate spatial footprints retain positional uncertainty.","Physical masks lack an independent accuracy sample or formal confusion matrix.","Spatial thinning disperses pixels but does not eliminate autocorrelation; tests remain exploratory.","No heat-index, physiological heat-stress or causal risk model has been implemented."])

    add_heading(doc,"6. Conclusion and Future Work",1)
    add_body(doc,"This project developed and validated a user-operated v0.3.2 workflow for apparent surface-temperature and Delta T analysis. Its scope extends beyond the five historical pilot images: users may select new inputs, make explicit correspondence and target decisions, generate native canonical results, reuse accepted captures, and request guarded spatial or temporal analysis. Five normal-route images and a separate football-field polygon/temporal case demonstrate distinct system capabilities.")
    add_body(doc,"The benchmark produced reproducible full-population and spatially thinned outputs, while the football-field case showed that capture-level temporal summaries can be generated without fabricating pixel correspondence. The project does not yet deliver a standard heat index or worker-rest advisory. Future work should collect repeated registered observations, independent air temperature, humidity, wind and radiation; calibrate the thermal system and material emissivity; quantify mask and alignment accuracy; use hierarchical or spatial models; and validate a decision layer against the Hong Kong Heat Index and occupational guidance. This progression preserves the project's long-term heat-risk objective while keeping current claims scientifically supportable.")

    add_heading(doc,"References",1)
    refs=["Benjamini, Y., & Hochberg, Y. (1995). Controlling the false discovery rate: A practical and powerful approach to multiple testing. Journal of the Royal Statistical Society: Series B, 57(1), 289-300. https://doi.org/10.1111/j.2517-6161.1995.tb02031.x","Colomina, I., & Molina, P. (2014). Unmanned aerial systems for photogrammetry and remote sensing: A review. ISPRS Journal of Photogrammetry and Remote Sensing, 92, 79-97. https://doi.org/10.1016/j.isprsjprs.2014.02.013","Hong Kong Labour Department. (2026). Guidance Notes on Prevention of Heat Stroke at Work (3rd ed.). https://www.labour.gov.hk/common/public/oh/Heat_Stress_GN_en.pdf","Kruskal, W. H., & Wallis, W. A. (1952). Use of ranks in one-criterion variance analysis. Journal of the American Statistical Association, 47(260), 583-621. https://doi.org/10.1080/01621459.1952.10483441","Mann, H. B., & Whitney, D. R. (1947). On a test of whether one of two random variables is stochastically larger than the other. Annals of Mathematical Statistics, 18(1), 50-60. https://doi.org/10.1214/aoms/1177730491","Oke, T. R. (1982). The energetic basis of the urban heat island. Quarterly Journal of the Royal Meteorological Society, 108(455), 1-24. https://doi.org/10.1002/qj.49710845502","Planning Department, Hong Kong SAR Government. (2024). Land Utilisation in Hong Kong 2024 [geospatial dataset].","Scott, D. W. (1992). Multivariate Density Estimation: Theory, Practice, and Visualization. Wiley.","Stewart, I. D., & Oke, T. R. (2012). Local Climate Zones for urban temperature studies. Bulletin of the American Meteorological Society, 93(12), 1879-1900. https://doi.org/10.1175/BAMS-D-11-00019.1","Voogt, J. A., & Oke, T. R. (2003). Thermal remote sensing of urban climates. Remote Sensing of Environment, 86(3), 370-384. https://doi.org/10.1016/S0034-4257(03)00079-8","Weng, Q., Lu, D., & Schubring, J. (2004). Estimation of land surface temperature-vegetation abundance relationship for urban heat island studies. Remote Sensing of Environment, 89(4), 467-483. https://doi.org/10.1016/j.rse.2003.11.005"]
    for ref in refs:
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.3); p.paragraph_format.first_line_indent=Inches(-.3); font(p.add_run(ref),10)

    add_heading(doc,"Appendix A. Reproducibility and validation boundary",1)
    add_table(doc,["Item","Recorded value"],[["Workflow","v0.3.2; canonical schema 0.2.0"],["Five-image population","1,638,400 finite pixels"],["Sampling","One pixel per 8 by 8 tile; seed 20260715; 20 secondary seeds"],["Temporal unit","One accepted capture of one physical target"],["Local verification","Python 3.12.13; 111 passed; 6 controlled-local deselected"],["Remote CI","Pending at draft date"],["Scientific boundary","Apparent surface temperature and Delta T; no heat-index model"]],[2700,6660],"Table A1. Key reproducibility settings and current evidence boundary.")
    add_heading(doc,"Appendix B. Draft status",1)
    add_bullets(doc,["This is the independent full rewrite v102 and does not replace the 4 August draft or the separate section-by-section v2 document.","Replace student, supervisor and department placeholders before submission.","Confirm the required university template, word limit, declaration wording and citation style.","Complete physical-plausibility and ambient-provenance review before freezing numerical interpretation.","Update the remote CI statement if GitHub Actions runs before submission."])
    doc.core_properties.title="Development and Validation of a User-Operated UAV Thermal Analysis Workflow"
    doc.core_properties.subject="UROP full rewrite v102"
    doc.core_properties.comments="Independent full rewrite; original draft preserved."
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
